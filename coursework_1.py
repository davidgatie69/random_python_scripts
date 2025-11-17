"""
GROUP 5 OOP
MUWANGUZI DAVID JEREMIAH
KIMULI DAVID 
NSUBUGA BRIAN
NASSALI FELISTA SSEGAWA
STEPHEN CYRUS KALEMA


MAKERERE UNIVERSITY - COURSEWORK 1
CGPA Calculator System with Basic Calculator
Question 1, 2, 3 Implementation
"""

import os
from datetime import datetime
from docx import Document
from docx.shared import Inches

class CGPAStudentSystem:
    
    def __init__(self):
        self.grade_points = {'A': 5, 'B': 4, 'C': 3, 'D': 2, 'E': 1, 'F': 0}
        self.student_records_document = Document()
        self._setup_document_header()
    
    def _setup_document_header(self):
        self.student_records_document.add_heading('MAKERERE UNIVERSITY - STUDENT RECORDS SYSTEM', 0)
        self.student_records_document.add_heading('CGPA Calculator and Basic Calculator Results', 1)
        self.student_records_document.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        self.student_records_document.add_paragraph()
    
    def calculate_semester_cgpa(self, grades_list, credit_units_list):       
        total_grade_points = 0
        total_credit_units = 0
        
        for i in range(len(grades_list)):
            grade_point = self.grade_points.get(grades_list[i].upper(), 0)
            total_grade_points += grade_point * credit_units_list[i]
            total_credit_units += credit_units_list[i]
        
        if total_credit_units == 0:
            return 0.0
        return total_grade_points / total_credit_units
    
    def perform_arithmetic_operation(self, number1, number2, operation):
        try:
            if operation == '+':
                return number1 + number2
            elif operation == '-':
                return number1 - number2
            elif operation == '*':
                return number1 * number2
            elif operation == '/':
                if number2 == 0:
                    raise ZeroDivisionError("Cannot divide by zero")
                return number1 / number2
            else:
                raise ValueError("Invalid operation")
        except Exception as error:
            return f"Error: {str(error)}"
    
    def classify_cgpa_performance(self, cgpa_value):
        if cgpa_value >= 4.4:
            return "Distinction"
        elif cgpa_value >= 3.5:
            return "Upper Credit"
        elif cgpa_value >= 2.0:
            return "Pass"
        else:
            return "Fail"
    
    def extract_student_number_values(self):
        student_numbers = [216022204, 216002204, 216007570, 216002774]
        total_sum = sum(student_numbers)
        
        # Remove the extreme value '8' on the left
        sum_string = str(total_sum)[1:]
        
        # Extract the four values
        value1 = int(sum_string[0:2])  # 64
        value2 = int(sum_string[2:4])  # 03
        value3 = int(sum_string[4:6])  # 47
        value4 = int(sum_string[6:8])  # 52
        
        return [value1, value2, value3, value4], total_sum
    
    def get_course_information_input(self):
        courses_list = []
        grades_list = []
        credit_units_list = []
        
        print("\n" + "="*50)
        print("CGPA CALCULATOR - COURSE INFORMATION INPUT")
        print("="*50)
        print("Enter details for 4 courses:")
        
        for course_number in range(4):
            print(f"\nCourse {course_number + 1}:")
            course_name = input("Enter course name: ")
            grade = input("Enter grade (A-F): ").upper()
            credit_unit = int(input("Enter credit unit: "))
            
            courses_list.append(course_name)
            grades_list.append(grade)
            credit_units_list.append(credit_unit)
        
        return courses_list, grades_list, credit_units_list
    
    def display_cgpa_results(self, courses_list, grades_list, credit_units_list, cgpa_value, performance_classification):
        print("\n" + "="*60)
        print("CGPA CALCULATION RESULTS")
        print("="*60)
        
        # Write to Word document
        self.student_records_document.add_heading('CGPA Calculation Results', 2)
        
        # Create results table
        results_table = self.student_records_document.add_table(rows=1, cols=4)
        results_table.style = 'Light Grid Accent 1'
        
        # Add table headers
        header_cells = results_table.rows[0].cells
        header_cells[0].text = 'Course Name'
        header_cells[1].text = 'Grade'
        header_cells[2].text = 'Credit Units'
        header_cells[3].text = 'Grade Points'
        
        # Add course data to table
        for i in range(len(courses_list)):
            row_cells = results_table.add_row().cells
            row_cells[0].text = courses_list[i]
            row_cells[1].text = grades_list[i]
            row_cells[2].text = str(credit_units_list[i])
            row_cells[3].text = str(self.grade_points[grades_list[i].upper()])
            
            print(f"{courses_list[i]}: Grade {grades_list[i]}, Credit Units {credit_units_list[i]}")
        
        # Add CGPA results
        self.student_records_document.add_paragraph(f"Calculated CGPA: {cgpa_value:.2f}")
        self.student_records_document.add_paragraph(f"Performance Classification: {performance_classification}")
        self.student_records_document.add_paragraph()
        
        print(f"\nCalculated CGPA: {cgpa_value:.2f}")
        print(f"Performance Classification: {performance_classification}")
        print("="*60)
    
    def save_student_semester_record(self, student_id, courses_list, grades_list, credit_units_list, cgpa_value, semester):
        self.student_records_document.add_heading(f'Student Academic Record - {student_id}', 3)
        self.student_records_document.add_paragraph(f"Student ID: {student_id}")
        self.student_records_document.add_paragraph(f"Semester: {semester}")
        self.student_records_document.add_paragraph(f"CGPA: {cgpa_value:.2f}")
        self.student_records_document.add_paragraph(f"Performance: {self.classify_cgpa_performance(cgpa_value)}")
        self.student_records_document.add_paragraph(f"Record Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        self.student_records_document.add_paragraph()
    
    def save_student_records_document(self):
        try:
            filename = f"Student_Records_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            self.student_records_document.save(filename)
            print(f"Student records document saved: {filename}")
            return filename
        except Exception as error:
            print(f"Error saving document: {error}")
            return None

def demonstrate_question_one():
    student_system = CGPAStudentSystem()
    
    print("\n" + "="*70)
    print("QUESTION 1: CGPA CALCULATOR AND BASIC CALCULATOR SYSTEM")
    print("="*70)
    
    # Part A: CGPA Calculation
    print("\nPART A: CGPA CALCULATION")
    courses, grades, credit_units = student_system.get_course_information_input()
    
    cgpa = student_system.calculate_semester_cgpa(grades, credit_units)
    classification = student_system.classify_cgpa_performance(cgpa)
    
    student_system.display_cgpa_results(courses, grades, credit_units, cgpa, classification)
    
    # Save student record
    student_id = input("\nEnter student ID: ")
    semester = input("Enter semester: ")
    student_system.save_student_semester_record(student_id, courses, grades, credit_units, cgpa, semester)
    
    # Part B: Basic Calculator with student numbers
    print("\nPART B: BASIC CALCULATOR WITH STUDENT NUMBERS")
    extracted_values, total_sum = student_system.extract_student_number_values()
    
    # Write calculator results to document
    student_system.student_records_document.add_heading('Basic Calculator Operations', 2)
    student_system.student_records_document.add_paragraph(f"Student numbers sum: {total_sum}")
    student_system.student_records_document.add_paragraph(f"Extracted values: {extracted_values}")
    
    print(f"Student numbers sum: {total_sum}")
    print(f"Extracted values: {extracted_values}")
    
    val1, val2, val3, val4 = extracted_values
    
    student_system.student_records_document.add_paragraph("Basic Calculator Operations:")
    
    operations_list = [
        (val1, val2, '+'),
        (val3, val4, '-'),
        (val1, val3, '*'),
        (val2, val4, '/')
    ]
    
    for num1, num2, operation in operations_list:
        result = student_system.perform_arithmetic_operation(num1, num2, operation)
        operation_string = f"{num1} {operation} {num2} = {result}"
        student_system.student_records_document.add_paragraph(operation_string)
        print(operation_string)
    
    # Save the document
    document_filename = student_system.save_student_records_document()
    
    return student_system, cgpa, classification, extracted_values, document_filename

# Question Two C Implementation
QUESTION_TWO_C_CODE = '''
#include <stdio.h>
#include <stdlib.h>

struct Course {
    char* courseName;
    int courseCode;
} courseInstance;

// Function prototype for weThink function
void weThink(struct Course courseData);

int main() {
    // a) Initialize the members of the structure
    struct Course courseInstance = {"Structured Programming", 101};
    
    // b) Allow user to enter data through keyboard
    printf("Enter course name: ");
    char userName[100];
    scanf("%99s", userName);
    courseInstance.courseName = userName;
    
    printf("Enter course code: ");
    scanf("%d", &courseInstance.courseCode);
    
    // c) Declare and initialize pointer to structure
    struct Course *coursePointer = &courseInstance;
    printf("Course Code accessed via pointer: %d\\n", coursePointer->courseCode);
    
    // d) Pass structure to function
    weThink(courseInstance);
    
    return 0;
}

// Function definition for weThink
void weThink(struct Course courseData) {
    printf("Course Name: %s, Course Code: %d\\n", courseData.courseName, courseData.courseCode);
}
'''

def create_question_two_document():
    """Create Word document for Question Two C code solution"""
    question_two_document = Document()
    question_two_document.add_heading('QUESTION TWO - C PROGRAMMING STRUCTURES', 0)
    question_two_document.add_heading('Structure Implementation and Pointer Operations', 1)
    question_two_document.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    question_two_document.add_paragraph()
    
    question_two_document.add_heading('C Code Implementation', 2)
    question_two_document.add_paragraph(QUESTION_TWO_C_CODE, style='Normal')
    
    question_two_document.add_heading('Expected Output Explanation', 2)
    question_two_document.add_paragraph('When executed, this C program will:')
    question_two_document.add_paragraph('• Initialize a course structure with default values', style='List Bullet')
    question_two_document.add_paragraph('• Accept user input for course name and code', style='List Bullet')
    question_two_document.add_paragraph('• Demonstrate pointer access to structure members', style='List Bullet')
    question_two_document.add_paragraph('• Pass structure to function by value', style='List Bullet')
    
    filename = "Question_Two_C_Implementation.docx"
    question_two_document.save(filename)
    print(f"Question Two document saved: {filename}")
    return filename

def demonstrate_question_three():
    question_three_document = Document()
    question_three_document.add_heading('QUESTION THREE - ERROR HANDLING AND OOP CONCEPTS', 0)
    question_three_document.add_heading('Python Exception Handling and Object-Oriented Programming', 1)
    question_three_document.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    question_three_document.add_paragraph()
    
    # Part A: Payment processing with error handling
    question_three_document.add_heading('Part A: Payment Processing API Error Handling', 2)
    
    def calculate_average_payment(amount_collected, number_of_users):
        """Calculate average amount per user with error handling"""
        try:
            if number_of_users == 0:
                raise ZeroDivisionError("Cannot divide by zero")
            return amount_collected / number_of_users
        except ZeroDivisionError:
            return "Error: Cannot divide by zero"
    
    # Test cases for payment processing
    test_cases = [(1000, 5), (1000, 0), (2500, 10)]
    question_three_document.add_paragraph("Payment Processing Test Cases:")
    
    for amount, users in test_cases:
        result = calculate_average_payment(amount, users)
        question_three_document.add_paragraph(f"Average (${amount}, {users} users): {result}")
    
    # Part B: File handling with error handling
    question_three_document.add_heading('Part B: Data-logging System File Handling', 2)
    
    def handle_daily_report_file():
        try:
            with open("daily_report.txt", "r") as file:
                content = file.read()
                return "File read successfully"
        except FileNotFoundError:
            return "File not found"
    
    file_result = handle_daily_report_file()
    question_three_document.add_paragraph(f"File handling result: {file_result}")
    
    # Part C: OOP Concepts with Inventory System
    question_three_document.add_heading('Part C: OOP Concepts - Inventory Management System', 2)
    
    # OOP Concepts Explanation
    question_three_document.add_heading('i) Object-Oriented Programming Concepts', 3)
    
    # Inheritance Explanation
    question_three_document.add_heading('Inheritance', 4)
    question_three_document.add_paragraph('''
    Inheritance allows a class to inherit attributes and methods from another class, 
promoting code reusability and establishing hierarchical relationships.

Example:
class Product:
    def __init__(self, name, price):
        self.name = name
        self.price = price

class Electronics(Product):
    def __init__(self, name, price, warranty):
        super().__init__(name, price)
        self.warranty_period = warranty
                                          ''')
    # Encapsulation Explanation
    question_three_document.add_heading('Encapsulation', 4)
    question_three_document.add_paragraph('''
    Encapsulation restricts direct access to some components and prevents
accidental modification of data through access modifiers.

Example:
                                          class BankAccount:
    def __init__(self, balance):
        self.__balance = balance  # Private attribute
    
    def get_balance(self):
        return self.__balance
    
    def deposit(self, amount):
        if amount > 0:
            self.__balance += amount
                                          ''')
    
    # Polymorphism Explanation
    question_three_document.add_heading('Polymorphism', 4)
    question_three_document.add_paragraph('''
    Polymorphism allows methods to have different implementations based on
the object they are acting upon, enabling flexible and extensible code.

Example:
class Shape:
    def area(self):
        pass

class Circle(Shape):
    def area(self, radius):
        return 3.14 * radius * radius

class Square(Shape):
    def area(self, side):
        return side * side
                                          ''')
    
    # Polymorphism Implementation Examples
    question_three_document.add_heading('ii) Polymorphism Implementation Examples', 3)
    question_three_document.add_paragraph('''
    class InventoryProduct:
        """Base product class for inventory management"""
        
        def __init__(self, product_name, product_price, product_category, stock_quantity):
            self.name = product_name
            self.price = product_price
            self.category = product_category
            self.stock_quantity = stock_quantity
        
        def apply_discount(self, discount_percentage):
            """Apply discount to product price"""
            self.price *= (1 - discount_percentage / 100)
            return self.price
        
        def display_product_info(self):
            """Display product information"""
            return f"{self.name} - ${self.price:.2f} - {self.category} - Stock: {self.stock_quantity}"

    class ElectronicsProduct(InventoryProduct):
        """Electronics product with additional features"""
        
        def __init__(self, product_name, product_price, stock_quantity, warranty_period):
            super().__init__(product_name, product_price, "Electronics", stock_quantity)
            self.warranty_period = warranty_period  # in months
        
        def apply_discount(self, discount_percentage):
            """Electronics get an extra 5% discount"""
            return self.price * (1 - (discount_percentage + 5) / 100)
        
        def display_product_info(self):
            """Display electronics product information"""
            base_info = super().display_product_info()
            return f"{base_info} - Warranty: {self.warranty_period} months"

    class ClothingProduct(InventoryProduct):
        """Clothing product with size information"""
        
        def __init__(self, product_name, product_price, stock_quantity, clothing_size):
            super().__init__(product_name, product_price, "Clothing", stock_quantity)
            self.size = clothing_size
        
        def apply_discount(self, discount_percentage):
            """Clothing items have a minimum price constraint"""
            discounted_price = self.price * (1 - discount_percentage / 100)
            return max(discounted_price, 10)  # Minimum $10
        
        def display_product_info(self):
            """Display clothing product information"""
            base_info = super().display_product_info()
            return f"{base_info} - Size: {self.size}"
    
    # Demonstrate polymorphism
    question_three_document.add_paragraph("Polymorphism in Action - Applying 15% Discount:")

    products = [
        ElectronicsProduct("Gaming Laptop", 1200, 15, 24),
        ClothingProduct("Cotton T-Shirt", 30, 100, "Large"),
        InventoryProduct("Programming Book", 45, "Books", 50)
    ]

    for product in products:
        original_price = product.price
        discounted_price = product.apply_discount(15)
        product_info = f"{product.name}: ${original_price:.2f} → ${discounted_price:.2f}"
        question_three_document.add_paragraph(product_info)

    filename = "Question_Three_OOP_Concepts.docx"
    question_three_document.save(filename)
    print(f"Question Three document saved: {filename}")
    return filename
                                          
        ''')

def main():
    print("MAKERERE UNIVERSITY - COURSEWORK 1 EXECUTION")
    print("="*70)
    # Execute Question One
    system, cgpa, classification, values, doc_file = demonstrate_question_one()

    # Execute Question Two
    q2_file = create_question_two_document()

    # Execute Question Three
    q3_file = demonstrate_question_three()

    print("\n" + "="*70)
    print("COURSEWORK 1 EXECUTION COMPLETED SUCCESSFULLY!")
    print("="*70)
    print(f"Question 1 Output: {doc_file}")
    print(f"Question 2 Output: {q2_file}")
    print(f"Question 3 Output: {q3_file}")
    print("="*70)

if __name__ == '__main__':
    main()